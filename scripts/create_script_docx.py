"""
Generates a professional Word (.docx) document containing the exact pitch script
for Scenes 15, 16, and 17 (2:45 – 3:30).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_script_docx(output_path):
    doc = docx.Document()
    
    # Page Setup (Letter, 1 inch margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x2B, 0x3A)
    
    # ------------------ TITLE & HEADER ------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("NASA SPACE APPS CHALLENGE 2026 // PRODUCTION SCRIPT\n")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x82, 0x99)  # Deep Cyan
    
    run_title = title_p.add_run("H.E.L.I.O.S. — 4-MINUTE PITCH SCRIPT")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0B, 0x1A, 0x30)  # Deep Aerospace Navy
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_part = subtitle_p.add_run("FINAL ACT: EXTREME ISOLATION, SCALABILITY & RESOLUTION (2:45 – 3:30)\n")
    run_part.font.name = 'Segoe UI'
    run_part.font.size = Pt(12)
    run_part.font.bold = True
    run_part.font.color.rgb = RGBColor(0x1B, 0x6C, 0x48)  # Emerald
    
    run_meta = subtitle_p.add_run("Team: Cosmic Plus  |  System: Health Evaluation Logistic Intelligent Onboard System  |  Deliverable: Video Scenes 15, 16, 17")
    run_meta.font.name = 'Segoe UI'
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)
    
    # Horizontal rule
    hr_table = doc.add_table(rows=1, cols=1)
    hr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hr_cell = hr_table.cell(0, 0)
    set_cell_background(hr_cell, "008299")
    hr_cell.width = Inches(6.5)
    set_cell_margins(hr_cell, top=10, bottom=10, left=0, right=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Helper for Scene Header
    def add_scene_header(scene_num, time_range, scene_title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(f"SCENE {scene_num} ")
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0x00, 0x82, 0x99)
        
        r2 = p.add_run(f"| {time_range} | ")
        r2.font.name = 'Segoe UI'
        r2.font.size = Pt(13)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)
        
        r3 = p.add_run(scene_title)
        r3.font.name = 'Segoe UI'
        r3.font.size = Pt(15)
        r3.font.bold = True
        r3.font.color.rgb = RGBColor(0x0B, 0x1A, 0x30)
        
    def add_voiceover_box(vo_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F0F5FA")  # Soft ice blue
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run("🎙️ VOICEOVER\n")
        r_lbl.font.name = 'Segoe UI'
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(0x00, 0x82, 0x99)
        
        r_txt = p.add_run(vo_text)
        r_txt.font.name = 'Georgia'
        r_txt.font.size = Pt(11)
        r_txt.font.italic = True
        r_txt.font.color.rgb = RGBColor(0x15, 0x25, 0x3C)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_visual_prompt(prompt_id, video_time, core_title, bullets, on_screen_text=None, visual_style=None):
        vp_p = doc.add_paragraph()
        vp_p.paragraph_format.space_before = Pt(10)
        vp_p.paragraph_format.space_after = Pt(4)
        
        r_vp = vp_p.add_run(f"VISUAL PROMPT {prompt_id} ")
        r_vp.font.name = 'Segoe UI'
        r_vp.font.size = Pt(11.5)
        r_vp.font.bold = True
        r_vp.font.color.rgb = RGBColor(0x1B, 0x6C, 0x48)  # Green
        
        r_time = vp_p.add_run(f"| VIDEO TIME: {video_time}")
        r_time.font.name = 'Segoe UI'
        r_time.font.size = Pt(10.5)
        r_time.font.bold = True
        r_time.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)
        
        lead_p = doc.add_paragraph()
        lead_p.paragraph_format.space_after = Pt(4)
        r_lead = lead_p.add_run(core_title)
        r_lead.font.name = 'Segoe UI'
        r_lead.font.size = Pt(10.5)
        r_lead.font.bold = True
        r_lead.font.color.rgb = RGBColor(0x0B, 0x1A, 0x30)
        
        for b in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.space_before = Pt(0)
            r = bp.add_run(b)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x22, 0x2B, 0x3A)
            
        if on_screen_text:
            ost_tbl = doc.add_table(rows=1, cols=1)
            ost_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            ost_cell = ost_tbl.cell(0, 0)
            ost_cell.width = Inches(6.5)
            set_cell_background(ost_cell, "F5F7FA")
            set_cell_margins(ost_cell, top=100, bottom=100, left=150, right=150)
            
            ost_p = ost_cell.paragraphs[0]
            ost_p.paragraph_format.space_after = Pt(0)
            r_ost_lbl = ost_p.add_run("ON-SCREEN UI / TEXT:  ")
            r_ost_lbl.font.name = 'Segoe UI'
            r_ost_lbl.font.size = Pt(9)
            r_ost_lbl.font.bold = True
            r_ost_lbl.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)
            
            r_ost_val = ost_p.add_run(on_screen_text)
            r_ost_val.font.name = 'Consolas'
            r_ost_val.font.size = Pt(9.5)
            r_ost_val.font.bold = True
            r_ost_val.font.color.rgb = RGBColor(0x0B, 0x1A, 0x30)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            
        if visual_style:
            vs_p = doc.add_paragraph()
            vs_p.paragraph_format.space_before = Pt(2)
            vs_p.paragraph_format.space_after = Pt(6)
            r_vs_lbl = vs_p.add_run("Visual Style: ")
            r_vs_lbl.font.name = 'Segoe UI'
            r_vs_lbl.font.size = Pt(9.5)
            r_vs_lbl.font.italic = True
            r_vs_lbl.font.bold = True
            r_vs_lbl.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)
            
            r_vs_val = vs_p.add_run(visual_style)
            r_vs_val.font.name = 'Segoe UI'
            r_vs_val.font.size = Pt(9.5)
            r_vs_val.font.italic = True
            r_vs_val.font.color.rgb = RGBColor(0x60, 0x72, 0x8A)

    # ==================== SCENE 15 ====================
    add_scene_header("15", "2:45–3:00", "EXTREME ISOLATION")
    add_voiceover_box(
        "“The true test of space medicine is not just deep space. It is extreme isolation.\n\n"
        "HELIOS works without cloud dependency, providing clinical decision support in Antarctic stations, "
        "disaster zones, and other environments where evacuation is difficult and communication can fail.”"
    )
    
    add_visual_prompt(
        "15A", "2:45–2:50",
        "Create a 5-second cinematic transition from deep-space medical monitoring to an isolated Antarctic research environment.",
        [
            "Begin with existing HELIOS astronaut telemetry interface showing live ECG waveform, heart rate, SpO₂, temperature, and respiratory data on a realistic aerospace console.",
            "Camera slowly pushes toward the telemetry data.",
            "Seamless transition from spacecraft environment into realistic Antarctic research station during severe polar blizzard (heavy snow, whipping wind, isolated modular buildings, limited visibility).",
            "Inside the station, showcase a rugged Panasonic Toughbook displaying the HELIOS interface, commercially realistic and portable.",
            "Core message: The same software intelligence operates in extreme terrestrial isolation."
        ],
        on_screen_text="EXTREME ISOLATION  |  LOCAL INTELLIGENCE  |  OFFLINE-FIRST CORE",
        visual_style="Cinematic documentary realism, realistic aerospace technology, cold Antarctic storm environment, professional NASA-style engineering presentation."
    )
    
    add_visual_prompt(
        "15B", "2:50–2:55",
        "Create a realistic 5-second disaster-medicine scene demonstrating offline clinical decision support.",
        [
            "Show an emergency medical responder working inside a temporary field clinic tent after a disaster with basic equipment and a rugged shockproof tablet.",
            "Tablet displays simplified HELIOS telemetry interface receiving practical physiological measurements (HEART RATE, SpO₂, TEMPERATURE, RESPIRATION).",
            "Show measurements entering HELIOS locally with prominent status indicator: LOCAL PROCESSING (no cloud server, no internet connection).",
            "System identifies concerning physiological trend (elevated HR, hypoxic drift, febrile elevation) and presents concise clinical decision-support advisory without replacing the physician."
        ],
        on_screen_text="TELEMETRY RECEIVED  →  LOCAL ANALYSIS  →  TREND DETECTED  →  CLINICAL ALERT",
        visual_style="Austere field hospital documentary realism, high-contrast clinical tablet HUD, non-autonomous decision support."
    )
    
    add_visual_prompt(
        "15C", "2:55–3:00",
        "Create a 5-second technical visualization that clearly communicates HELIOS as a reusable intelligence architecture rather than a specific piece of hardware.",
        [
            "Clean dark aerospace interface showcasing 3 environments arranged horizontally:",
            "  1. SPACE: Astronaut telemetry → HELIOS → Decision support",
            "  2. ANTARCTIC: Available field sensors → HELIOS → Decision support",
            "  3. DISASTER ZONE: Available medical measurements → HELIOS → Decision support",
            "Animate three data streams flowing continuously into the same central HELIOS intelligence layer.",
            "End with all three environments unified and synchronized to the same core intelligence.",
            "Core message: HELIOS does not require identical hardware everywhere; the architecture adapts to available data sources."
        ],
        on_screen_text="ONE INTELLIGENCE ARCHITECTURE  |  MANY EXTREME ENVIRONMENTS",
        visual_style="Professional aerospace systems architecture diagram, coordinate grid, glowing particle data bus."
    )
    
    # ==================== SCENE 16 ====================
    add_scene_header("16", "3:00–3:15", "AEROSPACE SCALABILITY")
    add_voiceover_box(
        "“Built for modular telemetry and edge intelligence, HELIOS can evolve with the mission.\n\n"
        "From commercial orbital crews to multi-crewed lunar habitats, new sensors and additional crew members "
        "can be integrated without rebuilding the core intelligence.”"
    )
    
    add_visual_prompt(
        "16A", "3:00–3:05",
        "Create a 5-second technical animation showing the modular architecture of HELIOS.",
        [
            "Central node labeled: HELIOS CORE (EDGE REASONING ENGINE).",
            "Surrounding radial telemetry inputs: ECG, SpO₂, TEMPERATURE, RESPIRATION, ENVIRONMENT, LOCATION.",
            "Animate clean high-frequency data lines flowing from each input into HELIOS Core.",
            "Introduce additional sensor module: [+ NEW SENSOR] hot-docking seamlessly to the central bus without changing the core architecture."
        ],
        on_screen_text="MODULAR TELEMETRY  |  EDGE INTELLIGENCE  |  EXTENSIBLE ARCHITECTURE",
        visual_style="Aerospace systems architecture diagram, vector bus lines, particle packets, zero architectural rework."
    )
    
    add_visual_prompt(
        "16B", "3:05–3:10",
        "Create a 5-second realistic HELIOS dashboard scaling from one crew member to a multi-crew mission.",
        [
            "Begin with single astronaut telemetry card: CREW 01 (CDR E. REID) with live ECG, HR, SpO₂, temp, and nominal status.",
            "Smooth camera pull-back and grid interpolation: dashboard smoothly expands to add CREW 02, CREW 03, CREW 04.",
            "Reorganizes automatically into a synchronized 4-crew monitoring grid.",
            "All 4 crew members display synchronized real-time physiological curves without interface clutter.",
            "Core message: ONE CREW → MULTIPLE CREW MEMBERS → SAME CORE SYSTEM."
        ],
        on_screen_text="ONE CREW  →  MULTIPLE CREW MEMBERS  →  SAME CORE SYSTEM",
        visual_style="Real-time multi-channel medical HUD, linear parallel scalability, tabular telemetry alignment."
    )
    
    add_visual_prompt(
        "16C", "3:10–3:15",
        "Create a 5-second cinematic transition from commercial orbital operations to future lunar operations.",
        [
            "Begin with realistic commercial orbital spacecraft interior containing several crew members with subtle overlaid HELIOS monitoring interface.",
            "Match-cut transition into realistic lunar-orbit / lunar-habitat environment (Artemis Base Camp).",
            "Show same HELIOS architecture monitoring multiple crew members and habitat ECLSS telemetry.",
            "Conclude on clean system architecture view: CREW + SENSORS → HELIOS → MISSION AWARENESS."
        ],
        on_screen_text="ADAPTABLE  |  SCALABLE  |  ARTEMIS AND BEYOND",
        visual_style="Cinematic spacecraft match-cut, Earth-to-Moon operational continuum, NASA Artemis branding."
    )
    
    # ==================== SCENE 17 ====================
    add_scene_header("17", "3:15–3:30", "CINEMATIC CALLBACK & CLOSE")
    add_voiceover_box(
        "“A heartbeat changes nearly four hundred thousand kilometers from Earth.”\n\n"
        "[0.8-second pause]\n\n"
        "“HELIOS turns that change into understanding, so the crew can act before a signal becomes a crisis.\n\n"
        "Monitor. Understand. Act.\n\n"
        "We are Team Cosmic Plus.”"
    )
    
    add_visual_prompt(
        "17A", "3:15–3:19",
        "Create a 4-second cinematic callback that precisely echoes the opening sequence of the HELIOS video.",
        [
            "Return to the same dark visual environment used in the opening hook.",
            "Single live ECG Lead-II waveform occupying the center of the frame.",
            "Waveform begins with stable, regular rhythm (72 bpm), then introduces a subtle but clearly detectable heartbeat change (ectopic irregularity & compensatory pause).",
            "Holds briefly with minimal ambient sound and a single clean medical monitor ECG tone.",
            "No unnecessary text. Pure cinematic focus: A heartbeat changes."
        ],
        on_screen_text="LEAD-II // 1.0 mV/cm // REAL-TIME",
        visual_style="Ultra-clean dark medical telemetry, emerald-to-amber waveform sweep, acoustic monitor beeps."
    )
    
    add_visual_prompt(
        "17B", "3:19–3:24",
        "Create a 5-second close-up visualization showing HELIOS converting the changing ECG signal into actionable understanding.",
        [
            "ECG waveform enters the HELIOS analysis interface.",
            "Left pane displays raw physiological telemetry (Lead-II voltage stream, Z-score dispersion +2.82σ).",
            "Right pane displays 5-stage sequential processing flow:",
            "  1. SIGNAL (10 Hz Bio-Sensor Telemetry Ingestion)",
            "  2. ANALYSIS (Sentry Matrix Variance & Waveform Morphology)",
            "  3. CONTEXT (Personal Baseline & Environmental Correlator)",
            "  4. RISK DETECTED (Premature Ventricular Contraction // ARF 0.68)",
            "  5. RECOMMENDED RESPONSE (Electrolyte review & 12-lead ECG review)",
            "Visual distinction between raw telemetry and interpreted clinical understanding."
        ],
        on_screen_text="RAW DATA  →  UNDERSTANDING  →  ACTION",
        visual_style="Sequential pipeline animation, clinical reasoning matrix, non-prescriptive decision support."
    )
    
    add_visual_prompt(
        "17C", "3:24–3:27",
        "Create a 3-second resolution shot where the HELIOS interface stabilizes after detecting the physiological change.",
        [
            "ECG waveform continues normally and regularly in nominal green.",
            "Analysis panel displays concise status summary: PHYSIOLOGICAL CHANGE DETECTED, CONTEXT ANALYZED, RESPONSE GUIDANCE AVAILABLE.",
            "Surrounding interface elements slowly and gracefully fade out until only the glowing central HELIOS identity remains."
        ],
        on_screen_text="PHYSIOLOGICAL CHANGE DETECTED  |  CONTEXT ANALYZED  |  RESPONSE GUIDANCE AVAILABLE",
        visual_style="Smooth visual stabilization, graceful telemetry fade, central identity focus."
    )
    
    add_visual_prompt(
        "17D", "3:27–3:30",
        "Create a clean 3-second final title card for the HELIOS project.",
        [
            "Fade completely to dark obsidian background.",
            "Reveal HELIOS logo in center.",
            "Below it, reveal project motto: MONITOR. UNDERSTAND. ACT.",
            "Reveal team identity: TEAM COSMIC PLUS // NASA SPACE APPS CHALLENGE 2026.",
            "Minimal, professional, NASA-standard typography with clean hold until the final frame."
        ],
        on_screen_text="HELIOS  |  MONITOR. UNDERSTAND. ACT.  |  TEAM COSMIC PLUS",
        visual_style="Minimalist title card, NASA aerospace typography, crisp letter-spacing."
    )
    
    # Final Title Summary Box
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    final_tbl = doc.add_table(rows=1, cols=1)
    final_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    final_cell = final_tbl.cell(0, 0)
    final_cell.width = Inches(6.5)
    set_cell_background(final_cell, "0B1A30")
    set_cell_margins(final_cell, top=200, bottom=200, left=240, right=240)
    
    fp = final_cell.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(4)
    
    rf1 = fp.add_run("H.E.L.I.O.S.\n")
    rf1.font.name = 'Segoe UI'
    rf1.font.size = Pt(24)
    rf1.font.bold = True
    rf1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    rf2 = fp.add_run("MONITOR.  UNDERSTAND.  ACT.\n\n")
    rf2.font.name = 'Segoe UI'
    rf2.font.size = Pt(13)
    rf2.font.bold = True
    rf2.font.color.rgb = RGBColor(0x00, 0xE6, 0x76)
    
    rf3 = fp.add_run("TEAM COSMIC PLUS\n")
    rf3.font.name = 'Segoe UI'
    rf3.font.size = Pt(14)
    rf3.font.bold = True
    rf3.font.color.rgb = RGBColor(0x00, 0xF0, 0xFF)
    
    rf4 = fp.add_run("NASA SPACE APPS CHALLENGE 2026")
    rf4.font.name = 'Segoe UI'
    rf4.font.size = Pt(10)
    rf4.font.color.rgb = RGBColor(0xA0, 0xB5, 0xD0)
    
    # ------------------ VIDEO ASSET MAPPING TABLE ------------------
    doc.add_page_break()
    p_map = doc.add_paragraph()
    p_map.paragraph_format.space_before = Pt(12)
    p_map.paragraph_format.space_after = Pt(8)
    r_map = p_map.add_run("PRODUCTION DELIVERABLE MAPPING // RENDERED VIDEO FILES")
    r_map.font.name = 'Segoe UI'
    r_map.font.size = Pt(14)
    r_map.font.bold = True
    r_map.font.color.rgb = RGBColor(0x0B, 0x1A, 0x30)
    
    map_table = doc.add_table(rows=1, cols=4)
    map_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = map_table.rows[0].cells
    hdr_titles = ["Scene ID", "Timeline", "Video Filename (1080p 60fps)", "Key Visual Core"]
    hdr_widths = [Inches(1.0), Inches(1.1), Inches(2.8), Inches(1.6)]
    
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].width = hdr_widths[i]
        set_cell_background(hdr_cells[i], "142238")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    deliverables = [
        ("Scene 15A", "2:45–2:50", "Scene_15A_02-45_02-50_Extreme_Isolation.mp4", "Antarctic blizzard & rugged laptop"),
        ("Scene 15B", "2:50–2:55", "Scene_15B_02-50_02-55_Disaster_Medicine.mp4", "Disaster triage tablet & alert"),
        ("Scene 15C", "2:55–3:00", "Scene_15C_02-55_03-00_One_Intelligence_Architecture.mp4", "3-environment universal architecture"),
        ("Scene 16A", "3:00–3:05", "Scene_16A_03-00_03-05_Modular_Telemetry.mp4", "Central core & hot-plug sensor"),
        ("Scene 16B", "3:05–3:10", "Scene_16B_03-05_03-10_Multi_Crew_Scaling.mp4", "1-crew to 4-crew synchronized grid"),
        ("Scene 16C", "3:10–3:15", "Scene_16C_03-10_03-15_Orbital_to_Lunar_Scale.mp4", "Orbital to Lunar base match-cut"),
        ("Scene 17A", "3:15–3:19", "Scene_17A_03-15_03-19_Cinematic_Callback_Heartbeat.mp4", "Lead-II rhythm change & ECG audio tone"),
        ("Scene 17B", "3:19–3:24", "Scene_17B_03-19_03-24_Signal_to_Action.mp4", "Raw telemetry vs. 5-step analysis"),
        ("Scene 17C", "3:24–3:27", "Scene_17C_03-24_03-27_Resolution_Stabilization.mp4", "Stabilized rhythm & logo focus"),
        ("Scene 17D", "3:27–3:30", "Scene_17D_03-27_03-30_Final_Title_Close.mp4", "Final title card & Team Cosmic Plus"),
        ("Master Cut", "2:45–3:30", "Scenes_15_to_17_02-45_03-30_Full_Sequence.mp4", "Complete 45s continuous broadcast cut")
    ]
    
    for row_idx, (sc_id, t_line, f_name, v_core) in enumerate(deliverables):
        row = map_table.add_row()
        bg_color = "F7F9FC" if row_idx % 2 == 0 else "FFFFFF"
        if sc_id == "Master Cut":
            bg_color = "EAF5EE"  # Light green for master cut
            
        cells = row.cells
        vals = [sc_id, t_line, f_name, v_core]
        for c_idx, val in enumerate(vals):
            cells[c_idx].width = hdr_widths[c_idx]
            set_cell_background(cells[c_idx], bg_color)
            set_cell_margins(cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Segoe UI' if c_idx != 2 else 'Consolas'
            r.font.size = Pt(9)
            if sc_id == "Master Cut":
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x6C, 0x48)
            else:
                r.font.color.rgb = RGBColor(0x22, 0x2B, 0x3A)
                
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Video_Script", "HELIOS_Pitch_Script_Scenes_15_to_17.docx")
    create_script_docx(out_file)
